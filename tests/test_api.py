import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from app.main import app

PHI_NOTE = (
    "Patient: Maria Gonzalez\n"
    "DOB: 04/12/1971\n"
    "MRN: 00451234\n"
    "Phone: (303) 555-0142\n"
    "Assessment: type 2 diabetes, continue metformin.\n"
)

CLEAN_TEXT = (
    "Quarterly revenue grew 12% year over year. The platform team shipped the new "
    "caching layer and reduced p95 latency by 40ms.\n"
)


@pytest.fixture(scope="module")
def client():
    # The context manager runs the lifespan handler, which loads the model.
    with TestClient(app) as c:
        yield c


def upload(client, name, content, content_type="application/octet-stream"):
    return client.post("/scan", files={"file": (name, content, content_type)})


def test_health(client):
    assert client.get("/health").json() == {"status": "ok"}


def test_clean_text_is_accepted(client):
    resp = upload(client, "report.txt", CLEAN_TEXT.encode())
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "accepted"
    assert body["contains_phi"] is False
    assert body["findings"] == []


def test_clinical_note_is_rejected(client):
    resp = upload(client, "note.txt", PHI_NOTE.encode())
    assert resp.status_code == 422
    body = resp.json()
    assert body["status"] == "rejected"
    for entity in ("PERSON", "DATE_OF_BIRTH", "MEDICAL_RECORD_NUMBER", "PHONE_NUMBER"):
        assert entity in body["entity_counts"], body["entity_counts"]


def test_response_does_not_echo_phi(client):
    resp = upload(client, "note.txt", PHI_NOTE.encode())
    assert "Gonzalez" not in resp.text
    assert "00451234" not in resp.text


@pytest.mark.parametrize(
    "text, entity",
    [
        ("SSN on file: 536-22-8726", "US_SSN"),
        ("Contact jane.doe@example.com for follow-up", "EMAIL_ADDRESS"),
        ("Member ID: XGH4482019 under the PPO plan", "HEALTH_PLAN_ID"),
        ("Medical record number 7712093 was merged", "MEDICAL_RECORD_NUMBER"),
    ],
)
def test_individual_identifiers_rejected(client, text, entity):
    resp = upload(client, "snippet.txt", text.encode())
    assert resp.status_code == 422
    assert entity in resp.json()["entity_counts"]


def test_csv_with_phi_is_rejected(client):
    csv = "name,ssn,diagnosis\nRobert Chen,536-22-8726,hypertension\n"
    resp = upload(client, "patients.csv", csv.encode(), "text/csv")
    assert resp.status_code == 422


def test_docx_with_phi_is_rejected(client):
    doc = Document()
    doc.add_paragraph("Discharge summary")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Patient"
    table.rows[0].cells[1].text = "Maria Gonzalez, MRN: 00451234"
    buf = io.BytesIO()
    doc.save(buf)
    resp = upload(client, "summary.docx", buf.getvalue())
    assert resp.status_code == 422
    assert "MEDICAL_RECORD_NUMBER" in resp.json()["entity_counts"]


def test_clean_docx_is_accepted(client):
    doc = Document()
    doc.add_paragraph(CLEAN_TEXT)
    buf = io.BytesIO()
    doc.save(buf)
    assert upload(client, "notes.docx", buf.getvalue()).status_code == 200


def test_image_only_pdf_is_refused(client):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    resp = upload(client, "scan.pdf", buf.getvalue(), "application/pdf")
    assert resp.status_code == 400


def test_corrupt_pdf_is_refused(client):
    assert upload(client, "broken.pdf", b"not really a pdf").status_code == 400


def test_unsupported_type(client):
    assert upload(client, "image.png", b"\x89PNG....").status_code == 415


def test_too_large(client):
    limit = client.app.state.settings.max_upload_bytes
    assert upload(client, "big.txt", b"a" * (limit + 1)).status_code == 413
