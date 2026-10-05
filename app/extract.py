"""Extract plain text, and any column headings, from uploaded files."""

import csv
import io
import json
from collections import Counter
from dataclasses import dataclass, field
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader

TEXT_EXTENSIONS = {".txt", ".csv", ".tsv", ".json", ".md", ".xml", ".html", ".htm", ".log", ".hl7", ".yaml", ".yml"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx"}


class UnsupportedFileType(Exception):
    pass


class ExtractionError(Exception):
    pass


@dataclass
class Extracted:
    text: str
    # Column heading (or JSON key) -> number of non-empty values under it.
    columns: Counter = field(default_factory=Counter)


def extract(filename: str, data: bytes) -> Extracted:
    ext = PurePath(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext or '(none)'}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    try:
        if ext == ".pdf":
            return Extracted(_pdf_text(data))
        if ext == ".docx":
            return _docx(data)
        text = _decode(data)
        if ext in (".csv", ".tsv"):
            return Extracted(text, _delimited_columns(text, "\t" if ext == ".tsv" else ","))
        if ext == ".json":
            return Extracted(text, _json_columns(text))
        return Extracted(text)
    except (UnsupportedFileType, ExtractionError):
        raise
    except Exception as exc:  # malformed PDFs/DOCX raise a variety of errors
        raise ExtractionError(f"Could not read {ext} file: {exc}") from exc


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if reader.pages and not text.strip():
        # Image-only (scanned) PDFs have no text layer; accepting them would
        # let PHI through unscanned, so refuse instead.
        raise ExtractionError("PDF has no extractable text (scanned image?); OCR is not supported")
    return text


def _docx(data: bytes) -> Extracted:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    columns: Counter = Counter()
    for table in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        parts.extend("\t".join(row) for row in rows)
        columns.update(_count_columns(rows))
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return Extracted("\n".join(parts), columns)


def _delimited_columns(text: str, delimiter: str) -> Counter:
    try:
        return _count_columns(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error:
        # Not well-formed CSV; the text is still scanned by Presidio.
        return Counter()


def _count_columns(rows) -> Counter:
    """Treat the first row as headings and count non-empty values under each."""
    rows = iter(rows)
    header = next(rows, None)
    counts: Counter = Counter({h.strip(): 0 for h in header or [] if h.strip()})
    for row in rows:
        for heading, value in zip(header, row):
            if heading.strip() and value.strip():
                counts[heading.strip()] += 1
    return counts


def _json_columns(text: str) -> Counter:
    try:
        doc = json.loads(text)
    except ValueError:
        return Counter()
    counts: Counter = Counter()
    stack = [doc]
    while stack:  # iterative, so deeply nested input can't hit the recursion limit
        node = stack.pop()
        if isinstance(node, dict):
            for key, value in node.items():
                if isinstance(value, (dict, list)):
                    stack.append(value)
                    filled = isinstance(value, list) and any(
                        not isinstance(v, (dict, list)) and _filled(v) for v in value
                    )
                else:
                    filled = _filled(value)
                counts[str(key)] += int(filled)
        elif isinstance(node, list):
            stack.extend(v for v in node if isinstance(v, (dict, list)))
    return counts


def _filled(value) -> bool:
    return value is not None and str(value).strip() != ""
